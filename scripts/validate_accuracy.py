"""Held-out categorisation accuracy check against a client's own template.

For each template given, the LAST few already-categorised rows are hidden
(blanked out in a throwaway copy, never touching your real file) and treated
as if they were a brand-new statement: the live pipeline - client history,
then learned corrections, then the HMRC-reference rule, then Claude - is
asked to categorise them, using only the client's remaining history. Each
answer is then compared with what the accountant had actually written
before it was hidden.

This answers "if we had processed this client's last N transactions fresh,
would we have gotten them right?" - using nothing but a client's own
existing working-papers file. It needs no separate corrections file, unlike
scripts/validate_feedback_accuracy.py, which checks against explicit
corrections recorded in Feedbacks/.

Costs real Anthropic API calls for every transaction not covered by
history, learned corrections, or the HMRC reference rule - same as
processing a real statement. Needs ANTHROPIC_API_KEY in .env; without it,
AI-decided rows come back 'Unknown', same as a real upload without a key.

Usage:
    venv\\Scripts\\python scripts\\validate_accuracy.py "path\\to\\template.xlsx"
    venv\\Scripts\\python scripts\\validate_accuracy.py "template1.xlsx" "template2.xlsx" --holdout 30
    venv\\Scripts\\python scripts\\validate_accuracy.py --folder "path\\to\\clients"
    venv\\Scripts\\python scripts\\validate_accuracy.py "template.xlsx" --json results.json
"""

import argparse
import contextlib
import io
import json
import shutil
import sys
import tempfile
from collections import Counter
from datetime import datetime
from pathlib import Path

ROOT = Path(__file__).parent.parent
sys.path.insert(0, str(ROOT))

import openpyxl  # noqa: E402
from dotenv import load_dotenv  # noqa: E402

load_dotenv(ROOT / '.env')

from core import categorise as cat  # noqa: E402
from core.columns import resolve_columns  # noqa: E402
from core.excel_writer import (  # noqa: E402
    _detect_sheet, extract_analysis_vocabulary, extract_client_examples,
)


def _num(v) -> float:
    return float(v) if isinstance(v, (int, float)) and not isinstance(v, bool) else 0.0


def _text(v) -> str:
    if v in (None, '') or (isinstance(v, str) and v.startswith('=')):
        return ''
    return str(v).strip()


def _find_templates(folder: Path) -> list[Path]:
    """Every .xlsx under folder that looks like a working-papers template:
    has a RAW-ish sheet with a recognised date, description and category
    column, and at least a handful of already-categorised rows. Skips
    processed outputs, Excel lock files, and files that are themselves bank
    statements (no category history of their own)."""
    found = []
    for path in sorted(folder.rglob('*.xlsx')):
        name = path.name
        if name.startswith('~$') or '_processed' in name.lower():
            continue
        try:
            wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
            raw_name = _detect_sheet(wb, 'raw')
            ws = wb[raw_name]
            rc = resolve_columns(ws)
            if not (rc['date'] and rc['desc'] and rc['uc']):
                continue
            categorised = 0
            for row in ws.iter_rows(min_row=2, max_row=min(ws.max_row, 500)):
                if row[rc['date'] - 1].value is not None and row[rc['uc'] - 1].value not in (None, ''):
                    categorised += 1
                    if categorised >= 10:
                        break
            wb.close()
            if categorised >= 10:
                found.append(path)
        except Exception:
            continue
    return found


def _load_rows(template_path: Path):
    """[(row_number, transaction_dict, true_category)] for every row with a
    date and a category, oldest first as they appear in the sheet."""
    wb = openpyxl.load_workbook(template_path, data_only=True)
    raw_name = _detect_sheet(wb, 'raw')
    ws = wb[raw_name]
    rc = resolve_columns(ws)
    if not (rc['date'] and rc['desc'] and rc['uc']):
        raise ValueError(f"Could not find Date / Description / category columns in {template_path.name}")

    rows = []
    for r in range(2, ws.max_row + 1):
        d = ws.cell(r, rc['date']).value
        cat_value = _text(ws.cell(r, rc['uc']).value)
        if d is None or not cat_value:
            continue
        money_in = _num(ws.cell(r, rc['in']).value) if rc['in'] else 0.0
        money_out = abs(_num(ws.cell(r, rc['out']).value)) if rc['out'] else 0.0
        amt = _num(ws.cell(r, rc['amt']).value) if rc['amt'] else (_num(ws.cell(r, rc['net']).value) if rc['net'] else 0.0)
        if not money_in and not money_out:
            money_in, money_out = (amt, 0.0) if amt > 0 else (0.0, -amt)
        if not money_in and not money_out:
            continue
        rows.append((r, {
            'date': d.date() if hasattr(d, 'date') else d,
            'description': _text(ws.cell(r, rc['desc']).value),
            'payee': _text(ws.cell(r, rc['name']).value) if rc['name'] else '',
            'subcategory': _text(ws.cell(r, rc['type']).value) if rc['type'] else '',
            'reference': _text(ws.cell(r, rc['ref']).value) if rc['ref'] else '',
            'csv_category': _text(ws.cell(r, rc['bank_cat']).value) if rc['bank_cat'] else '',
            'money_in': money_in, 'money_out': money_out,
        }, cat_value))
    return rows, rc, raw_name


def _same(pred: str, truth: str, canon) -> bool:
    return cat._cat_key(canon(pred)) == cat._cat_key(canon(truth)) or cat._cat_key(pred) == cat._cat_key(truth)


def validate(template_path: Path, holdout: int) -> dict:
    rows, rc, raw_name = _load_rows(template_path)
    if len(rows) < holdout + 10:
        return {'template': str(template_path), 'skipped': f'only {len(rows)} categorised rows (need at least {holdout + 10})'}
    test = rows[-holdout:]
    history_rows = rows[:-holdout]

    # A throwaway copy with the held-out rows' categories blanked out - the
    # real file is never opened for writing.
    tmp_dir = Path(tempfile.mkdtemp(prefix='plutus_validate_'))
    try:
        hist_path = tmp_dir / template_path.name
        shutil.copy(template_path, hist_path)
        wb = openpyxl.load_workbook(hist_path)
        ws = wb[raw_name]
        for r, _, _ in test:
            ws.cell(r, rc['uc']).value = None
            for c in rc['reviewed']:
                ws.cell(r, c).value = None
        wb.save(hist_path)

        buf = io.StringIO()
        with contextlib.redirect_stdout(buf):
            examples = extract_client_examples(hist_path)
            vocab = extract_analysis_vocabulary(hist_path)
            txns = [t for _, t, _ in test]
            predictions, review = cat.categorise_detailed(txns, client_examples=examples, analysis_vocab=vocab)

            # Replay step 1 of categorise_detailed to attribute each answer's
            # source (history / learned / HMRC reference / AI) for reporting.
            rules, ambiguous = cat._history_rules(examples)
            learned = cat._load_learned()
            sources = []
            for t in txns:
                rule_cat, src = cat._rule_for(t, rules, ambiguous, learned)
                if src != 'history' and cat._hmrc_by_reference(t):
                    src = 'reference'
                elif src is None:
                    src = 'ai'
                sources.append(src)
    finally:
        shutil.rmtree(tmp_dir, ignore_errors=True)

    canon = cat._make_canonicaliser([c for _, _, c in history_rows], vocab)
    results = []
    for (r, t, truth), pred, rv, src in zip(test, predictions, review, sources):
        results.append({
            'row': r, 'date': str(t['date']), 'description': t['description'][:60],
            'truth': truth, 'predicted': pred, 'source': src,
            'correct': _same(pred, truth, canon), 'flagged': bool(rv),
        })

    by_source = {}
    for src in ('history', 'learned', 'reference', 'ai'):
        xs = [x for x in results if x['source'] == src]
        if xs:
            by_source[src] = {'n': len(xs), 'correct': sum(x['correct'] for x in xs)}

    return {
        'template': str(template_path),
        'held_out_rows': len(results),
        'correct': sum(x['correct'] for x in results),
        'flagged': sum(x['flagged'] for x in results),
        'wrong_and_not_flagged': sum(1 for x in results if not x['correct'] and not x['flagged']),
        'by_source': by_source,
        'rows': results,
    }


def _print_report(result: dict) -> None:
    name = Path(result['template']).name
    if 'skipped' in result:
        print(f"[skip] {name}: {result['skipped']}")
        return
    n = result['held_out_rows']
    ok = result['correct']
    print(f"{name}: {ok}/{n} correct ({100 * ok / n:.0f}%), "
          f"{result['flagged']} flagged for review, "
          f"{result['wrong_and_not_flagged']} wrong and NOT flagged")
    for src, s in result['by_source'].items():
        print(f"    {src:10} {s['correct']}/{s['n']} = {100 * s['correct'] / s['n']:.0f}%")
    for row in result['rows']:
        if not row['correct']:
            flag = '?' if row['flagged'] else ' '
            print(f"    {row['source']:10}{flag} said {row['predicted']!r:24} "
                  f"accountant said {row['truth']!r:24} | {row['description']}")


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument('templates', nargs='*', help='Path(s) to working-papers .xlsx template(s)')
    parser.add_argument('--folder', help='Scan this folder (and subfolders) for templates automatically')
    parser.add_argument('--holdout', type=int, default=50, help='How many of the most recent rows to hide and test (default: 50)')
    parser.add_argument('--json', help='Also write full per-row results to this JSON file')
    args = parser.parse_args()

    paths = [Path(p) for p in args.templates]
    if args.folder:
        found = _find_templates(Path(args.folder))
        print(f"Found {len(found)} template(s) under {args.folder}:")
        for p in found:
            print(f"  {p}")
        paths += found
    if not paths:
        parser.error('Give at least one template path, or --folder to scan one automatically.')

    api_key_set = bool(__import__('os').environ.get('ANTHROPIC_API_KEY', '')) and \
        not __import__('os').environ['ANTHROPIC_API_KEY'].startswith('sk-ant-...')
    if not api_key_set:
        print("WARNING: no ANTHROPIC_API_KEY set - every row the AI would decide "
              "will come back 'Unknown' (same as a real upload without a key).\n")

    all_results = []
    totals = Counter()
    print(f"Holding out the last {args.holdout} categorised row(s) of each template.\n")
    for path in paths:
        try:
            result = validate(path, args.holdout)
        except Exception as exc:
            result = {'template': str(path), 'skipped': f'{type(exc).__name__}: {exc}'}
        all_results.append(result)
        _print_report(result)
        print()
        if 'rows' in result:
            totals['n'] += result['held_out_rows']
            totals['ok'] += result['correct']
            totals['flagged'] += result['flagged']
            totals['wrong_unflagged'] += result['wrong_and_not_flagged']

    if totals['n']:
        print('=' * 70)
        print(f"OVERALL: {totals['ok']}/{totals['n']} correct "
              f"({100 * totals['ok'] / totals['n']:.1f}%)  |  "
              f"flagged {totals['flagged']} ({100 * totals['flagged'] / totals['n']:.1f}%)  |  "
              f"wrong-and-not-flagged {totals['wrong_unflagged']} "
              f"({100 * totals['wrong_unflagged'] / totals['n']:.1f}%)")

    if args.json:
        out = {'run_at': datetime.now().isoformat(timespec='seconds'), 'holdout': args.holdout, 'results': all_results}
        Path(args.json).write_text(json.dumps(out, indent=2, default=str), encoding='utf-8')
        print(f"\nFull per-row results written to {args.json}")


if __name__ == '__main__':
    main()

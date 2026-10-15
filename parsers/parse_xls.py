"""Universal bank statement spreadsheet parser (.xlsx and legacy .xls).

Reads each worksheet directly (keeping native date and number cell types,
rather than round-tripping everything through strings) and shares
column-detection / row-parsing with parse_csv.py (see parsers/_common.py),
so a spreadsheet and a CSV export of the same statement produce identical
output — handling separate Paid in / Paid out columns just as well as a
single signed Amount column (Barclays standard).

.xlsx is read with openpyxl; Excel 97-2003 .xls files with xlrd.
"""

from pathlib import Path

import openpyxl

from parsers._common import (DATE_HEADERS, find_col, find_header_row, fix_swapped_dates,
                             parse_rows, rows_to_dicts)


def _sheets_xlsx(path: Path) -> list:
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    try:
        return [list(ws.iter_rows(values_only=True)) for ws in wb.worksheets]
    finally:
        wb.close()


def _sheets_xls(path: Path) -> list:
    try:
        import xlrd
    except ImportError as exc:
        raise ValueError(
            "Reading .xls files needs the 'xlrd' package "
            "(pip install -r requirements.txt). Alternatively save the file "
            "as .xlsx or CSV."
        ) from exc

    try:
        book = xlrd.open_workbook(str(path))
    except xlrd.XLRDError as exc:
        raise ValueError(f"Could not open '{path.name}' as an Excel .xls file: {exc}") from exc
    try:
        sheets = []
        for sh in book.sheets():
            rows = []
            for r in range(sh.nrows):
                row = []
                for c in range(sh.ncols):
                    cell = sh.cell(r, c)
                    if cell.ctype in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK):
                        row.append(None)
                    elif cell.ctype == xlrd.XL_CELL_DATE:
                        try:
                            row.append(xlrd.xldate.xldate_as_datetime(cell.value, book.datemode))
                        except (xlrd.xldate.XLDateError, OverflowError, ValueError):
                            row.append(cell.value)
                    elif cell.ctype == xlrd.XL_CELL_BOOLEAN:
                        row.append(bool(cell.value))
                    elif cell.ctype == xlrd.XL_CELL_ERROR:
                        row.append(None)
                    else:
                        row.append(cell.value)
                rows.append(row)
            sheets.append(rows)
        return sheets
    finally:
        book.release_resources()


def parse_xls(path) -> list[dict]:
    path = Path(path)
    sheets = _sheets_xls(path) if path.suffix.lower() == '.xls' else _sheets_xlsx(path)

    # Use the first sheet that actually has a transaction header row —
    # exports sometimes lead with a summary / cover sheet.
    chosen = next((rows for rows in sheets if find_header_row(rows) is not None),
                  sheets[0] if sheets else [])
    headers, dict_rows = rows_to_dicts(chosen)
    if not dict_rows:
        return []
    fix_swapped_dates(dict_rows, find_col(headers, *DATE_HEADERS))
    return parse_rows(dict_rows, headers)

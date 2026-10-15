"""Shared fixtures. Every statement here is invented: made-up payees, round
amounts and no real names, account numbers or sort codes."""

import sys
from datetime import date
from pathlib import Path

import openpyxl
import pytest

sys.path.insert(0, str(Path(__file__).parent.parent))

# (date, description, type, paid in, paid out, balance after) — oldest first.
STATEMENT = [
    (date(2025, 4, 3),  'ACME OPTICAL SUPPLIES LTD',  'Faster payment', 0.00,    120.00, 880.00),
    (date(2025, 4, 5),  'EXAMPLE NHS PAYMENTS',       'BACS credit',    1500.00, 0.00,   2380.00),
    (date(2025, 4, 6),  'SHELL PETROL STATION',       'Card payment',   0.00,    45.50,  2334.50),
    (date(2025, 4, 6),  'SHELL PETROL STATION',       'Card payment',   0.00,    45.50,  2289.00),
    (date(2025, 5, 2),  'DEMO HIGH ST PARKING',       'Card payment',   0.00,    3.20,   2285.80),
]


def csv_text(rows, newest_first=False, header_block=False, signed=False) -> str:
    lines = []
    if header_block:
        lines += ['Account name: DEMO CLIENT', 'Statement period: Apr-May 2025', '']
    if signed:
        lines.append('Date,Description,Type,Amount,Balance')
    else:
        lines.append('Date,Description,Type,Paid in,Paid out,Balance')
    body = list(reversed(rows)) if newest_first else list(rows)
    for d, desc, typ, pin, pout, bal in body:
        ds = d.strftime('%d/%m/%Y')
        if signed:
            lines.append(f'{ds},{desc},{typ},{pin - pout:.2f},{bal:.2f}')
        else:
            lines.append(f'{ds},{desc},{typ},{pin or ""},{pout or ""},{bal:.2f}')
    return '\n'.join(lines) + '\n'


@pytest.fixture
def make_csv(tmp_path):
    def _make(name='statement.csv', **kw):
        p = tmp_path / name
        p.write_text(csv_text(STATEMENT, **kw), encoding='utf-8')
        return p
    return _make


@pytest.fixture
def make_xlsx(tmp_path):
    def _make(name='statement.xlsx', header_block=False, cover_sheet=False):
        wb = openpyxl.Workbook()
        ws = wb.active
        if cover_sheet:
            ws.title = 'Summary'
            ws.append(['Account summary'])
            ws = wb.create_sheet('Transactions')
        if header_block:
            ws.append(['Account name: DEMO CLIENT'])
            ws.append([])
        ws.append(['Date', 'Description', 'Type', 'Paid in', 'Paid out', 'Balance'])
        for d, desc, typ, pin, pout, bal in STATEMENT:
            ws.append([d, desc, typ, pin or None, pout or None, bal])   # real date cells
        p = tmp_path / name
        wb.save(p)
        return p
    return _make


@pytest.fixture
def make_xls(tmp_path):
    xlwt = pytest.importorskip('xlwt')

    def _make(name='statement.xls'):
        wb = xlwt.Workbook()
        ws = wb.add_sheet('Transactions')
        fmt = xlwt.easyxf(num_format_str='DD/MM/YYYY')
        for c, h in enumerate(['Date', 'Description', 'Type', 'Paid in', 'Paid out', 'Balance']):
            ws.write(0, c, h)
        for r, (d, desc, typ, pin, pout, bal) in enumerate(STATEMENT, start=1):
            ws.write(r, 0, d, fmt)
            ws.write(r, 1, desc)
            ws.write(r, 2, typ)
            if pin:
                ws.write(r, 3, pin)
            if pout:
                ws.write(r, 4, pout)
            ws.write(r, 5, bal)
        p = tmp_path / name
        wb.save(str(p))
        return p
    return _make


@pytest.fixture
def make_template(tmp_path):
    """A small working-papers template: a RAW tab with two existing rows
    (formulas included) and an Analysis tab."""
    def _make(name='template.xlsx'):
        wb = openpyxl.Workbook()
        raw = wb.active
        raw.title = 'RAW'
        raw.append(['No', 'SA', 'ACC', 'Date', 'Description', 'Paid in', 'Paid out',
                    'Balance', 'NET', 'UC Category'])
        raw.append([1, 'FY24', '24/25', date(2025, 3, 28), 'EXAMPLE NHS PAYMENTS',
                    1000.00, None, 1000.00, '=F2-G2', 'Income'])
        raw.append([2, 'FY24', '24/25', date(2025, 3, 30), 'ACME OPTICAL SUPPLIES LTD',
                    None, 0.00, 1000.00, '=F3-G3', 'Purchase'])
        an = wb.create_sheet('Analysis')
        an['A1'] = 'Category'
        an['B1'] = 'FY24'
        an['A2'] = 'Income'
        an['B2'] = '=SUMIFS(RAW!I:I,RAW!B:B,B$1,RAW!J:J,$A2)'
        p = tmp_path / name
        wb.save(p)
        return p
    return _make

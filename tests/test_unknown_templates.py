"""Templates and statements the program has never seen.

Every accountant names the columns differently, so the writer must work out
what each template column is for. These templates are invented, with
deliberately unfamiliar headers modelled on real client files.
"""

from datetime import date, timedelta

import openpyxl
import pytest

from conftest import STATEMENT
from core.columns import resolve_columns
from core.excel_writer import prepare_new_transactions, write_workbook
from parsers.parse_csv import parse_csv

PAYEES = [('EXAMPLE NHS PAYMENTS', 'in', 'BACS credit', 'INCOME', 'Income'),
          ('SHELL PETROL STATION', 'out', 'Card payment', 'TRAVEL', 'Petrol'),
          ('ACME OPTICAL SUPPLIES LTD', 'out', 'Faster payment', 'SUPPLIES', 'Purchase')]


def build_template(tmp_path, columns, name='t.xlsx', n_rows=40):
    """columns: list of (header, role) — role says what to put in history rows.
    Roles: no, date, year, year2, desc, ref, amt, in, out, bal, net, bank, uc,
    mirror, type, txid, time, blank."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'RAW'
    ws.append([h for h, _ in columns])
    bal = 100.0
    start = date(2024, 4, 8)
    for i in range(n_rows):
        d = start + timedelta(days=i * 6)
        payee, direction, typ, bank, uc = PAYEES[i % 3]
        amount = 50.0
        bal += amount if direction == 'in' else -amount
        year = 'FY25' if d < date(2025, 4, 6) else 'FY26'
        row = []
        for _, role in columns:
            row.append({
                'no': i + 1, 'date': d, 'year': year,
                'year2': '24/25' if year == 'FY25' else '25/26',
                'desc': payee, 'ref': f'REF{i}', 'type': typ,
                'amt': amount if direction == 'in' else -amount,
                'in': amount if direction == 'in' else None,
                'out': amount if direction == 'out' else None,
                'bal': round(bal, 2), 'net': None, 'bank': bank if False else bank,
                'uc': uc, 'mirror': uc, 'txid': f'tx_{i:05d}', 'time': '10:00:00',
                'blank': None,
            }[role])
        ws.append(row)
    # make the last balance 1000 so the invented statement reconciles
    bal_col = next((i for i, (_, r) in enumerate(columns, 1) if r == 'bal'), None)
    if bal_col:
        ws.cell(n_rows + 1, bal_col).value = 1000.00
    p = tmp_path / name
    wb.save(p)
    return p


def run(make_csv, template, tmp_path, cats=None):
    txns, _ = prepare_new_transactions(parse_csv(make_csv()), template)
    warnings = []
    out = tmp_path / 'out.xlsx'
    problems = write_workbook(txns, cats or ['Purchase', 'Income', 'Petrol', 'Petrol', 'Parking'],
                              template, out, warnings=warnings)
    ws = openpyxl.load_workbook(out)['RAW']
    return ws, txns, problems, warnings


HOWAYDA_LIKE = [('No', 'no'), ('Date', 'date'), ('Year', 'year'), ('Personal', 'year2'),
                ('Description', 'desc'), ('Description 2', 'ref'), ('Amount', 'amt'),
                ('Paid in', 'in'), ('Paid Out', 'out'), ('Balance', 'bal'),
                ('Spending', 'bank'), ('MUC desc', 'mirror'), ('UC desc', 'uc')]

HAROON_LIKE = [('NO', 'no'), ('Date', 'date'), ('Counter Party', 'desc'),
               ('Reference', 'ref'), ('Type', 'type'), ('Paid In', 'in'),
               ('Paid Out', 'out'), ('Balance (GBP)', 'bal'), ('Category Bank', 'bank'),
               ('Category  UC', 'uc'), ('SA', 'year'), ('ACC', 'year2')]

MONZO_LIKE = [('No', 'no'), ('Transaction ID', 'txid'), ('Date', 'date'), ('Time', 'time'),
              ('Type', 'type'), ('Description', 'desc'), ('Category', 'bank'),
              ('Money In', 'in'), ('Money Out', 'out'), ('Balance', 'bal'),
              ('UC category', 'uc')]


def test_column_roles_are_found_by_header_words_and_contents(tmp_path):
    ws = openpyxl.load_workbook(build_template(tmp_path, HOWAYDA_LIKE))['RAW']
    rc = resolve_columns(ws)
    headers = {c: ws.cell(1, c).value for c in range(1, ws.max_column + 1)}
    assert headers[rc['uc']] == 'UC desc'
    assert [headers[c] for c in rc['reviewed']] == ['MUC desc']
    assert headers[rc['bank_cat']] == 'Spending'
    assert headers[rc['ref']] == 'Description 2'
    # 'Year' (FY25) and 'Personal' (24/25) are recognised from their contents
    assert {headers[rc['acc']], headers[rc['fy']]} == {'Year', 'Personal'}


@pytest.mark.parametrize('columns', [HOWAYDA_LIKE, HAROON_LIKE])
def test_unfamiliar_headers_still_get_every_column_filled(make_csv, tmp_path, columns):
    template = build_template(tmp_path, columns)
    ws, txns, problems, warnings = run(make_csv, template, tmp_path)
    hdr = {ws.cell(1, c).value: c for c in range(1, ws.max_column + 1)}
    new_rows = range(42, 42 + len(txns))
    uc = hdr.get('UC desc') or hdr.get('Category  UC')
    assert [ws.cell(r, uc).value for r in new_rows] == ['Purchase', 'Income', 'Petrol', 'Petrol', 'Parking']
    year_cols = [hdr[h] for h in ('Year', 'Personal', 'SA', 'ACC') if h in hdr]
    for r in new_rows:
        assert all(ws.cell(r, c).value for c in year_cols), 'year labels missing'
    assert problems == []
    assert warnings == [] or all('Description 2' in w or 'Reference' in w or 'reviewer' in w
                                 for w in warnings)


def test_reviewed_column_is_left_for_the_reviewer(make_csv, tmp_path):
    ws, txns, _, warnings = run(make_csv, build_template(tmp_path, HOWAYDA_LIKE), tmp_path)
    hdr = {ws.cell(1, c).value: c for c in range(1, ws.max_column + 1)}
    for r in range(42, 42 + len(txns)):
        assert ws.cell(r, hdr['UC desc']).value            # the AI's first pass
        assert ws.cell(r, hdr['MUC desc']).value is None   # the reviewer's final answer
    assert any('reviewer' in w for w in warnings)


def test_statement_columns_matching_template_headers_are_copied(tmp_path):
    # Monzo-style: the template has Transaction ID / Time / Category columns
    csv = tmp_path / 'm.csv'
    csv.write_text(
        'Transaction ID,Date,Time,Type,Description,Category,Amount,Balance\n'
        'tx_900,03/04/2025,09:15:00,Card payment,SHELL PETROL STATION,Transport,-20.00,980.00\n',
        encoding='utf-8')
    template = build_template(tmp_path, MONZO_LIKE)
    txns, _ = prepare_new_transactions(parse_csv(csv), template)
    out = tmp_path / 'o.xlsx'
    write_workbook(txns, ['Petrol'], template, out)
    ws = openpyxl.load_workbook(out)['RAW']
    hdr = {ws.cell(1, c).value: c for c in range(1, ws.max_column + 1)}
    assert ws.cell(42, hdr['Transaction ID']).value == 'tx_900'
    assert ws.cell(42, hdr['Category']).value == 'Transport'      # the bank's, not the UC one
    assert ws.cell(42, hdr['UC category']).value == 'Petrol'
    assert ws.cell(42, hdr['Time']).value is not None


def test_payment_type_follows_the_clients_spelling_and_history(tmp_path):
    # a PDF-style statement: no Type, and 'Faster Payments' wording
    csv = tmp_path / 'p.csv'
    csv.write_text('Date,Description,Type,Paid in,Paid out,Balance\n'
                   '03/04/2025,ACME OPTICAL SUPPLIES LTD,Faster Payments,,20.00,980.00\n'
                   '04/04/2025,SHELL PETROL STATION,,,5.00,975.00\n', encoding='utf-8')
    template = build_template(tmp_path, HAROON_LIKE)
    txns, _ = prepare_new_transactions(parse_csv(csv), template)
    out = tmp_path / 'o.xlsx'
    write_workbook(txns, ['Purchase', 'Petrol'], template, out)
    ws = openpyxl.load_workbook(out)['RAW']
    hdr = {ws.cell(1, c).value: c for c in range(1, ws.max_column + 1)}
    assert ws.cell(42, hdr['Type']).value == 'Faster payment'      # client's spelling
    assert ws.cell(43, hdr['Type']).value == 'Card payment'        # learned from history
    assert ws.cell(43, hdr['Category Bank']).value == 'TRAVEL'     # learned, in their spelling


def test_a_template_column_nobody_can_fill_is_reported(make_csv, tmp_path):
    # unique IDs in history, nothing like it in the statement: cannot be guessed
    cols = HAROON_LIKE + [('Internal ID', 'txid')]
    ws, txns, _, warnings = run(make_csv, build_template(tmp_path, cols), tmp_path)
    assert ws.cell(42, len(cols)).value is None               # not invented
    assert any("Internal ID" in w for w in warnings)


def test_missing_category_column_is_reported(make_csv, tmp_path):
    cols = [c for c in HAROON_LIKE if c[1] not in ('uc',)]
    _, _, _, warnings = run(make_csv, build_template(tmp_path, cols), tmp_path)
    assert any('category column' in w for w in warnings)

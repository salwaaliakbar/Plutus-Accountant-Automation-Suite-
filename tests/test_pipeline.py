"""Statement -> template -> output workbook, without calling the AI."""

import openpyxl

from core import categorise as cat
from core.excel_writer import prepare_new_transactions, write_workbook
from parsers.parse_csv import parse_csv


def _run(make_csv, make_template, tmp_path, categories=None, review=None):
    template = make_template()
    txns, dropped = prepare_new_transactions(parse_csv(make_csv()), template)
    cats = categories or ['Purchase', 'Income', 'Petrol', 'Petrol', 'Parking']
    out = tmp_path / 'out.xlsx'
    problems = write_workbook(txns, cats, template, out, review=review)
    return template, out, txns, dropped, problems


def test_rows_appended_with_formulas_and_year_labels(make_csv, make_template, tmp_path):
    _, out, txns, _, problems = _run(make_csv, make_template, tmp_path)
    ws = openpyxl.load_workbook(out)['RAW']
    rows = list(ws.iter_rows(min_row=2, values_only=True))

    assert len(rows) == 2 + len(txns)              # 2 template rows + new rows
    assert [r[0] for r in rows] == list(range(1, len(rows) + 1))
    assert all(r[8] == f'=F{i}-G{i}' for i, r in enumerate(rows, start=2))  # NET formula copied
    # 3 Apr is still FY24 (UK year starts 6 Apr); 6 Apr onwards is FY25
    assert [r[1] for r in rows] == ['FY24'] * 4 + ['FY25'] * 3
    assert problems == []                          # balances reconcile


def test_original_rows_are_not_changed(make_csv, make_template, tmp_path):
    template, out, *_ = _run(make_csv, make_template, tmp_path)
    before = list(openpyxl.load_workbook(template)['RAW'].iter_rows(min_row=1, max_row=2, values_only=True))
    after = list(openpyxl.load_workbook(out)['RAW'].iter_rows(min_row=1, max_row=2, values_only=True))
    assert [r[:10] for r in after] == [r[:10] for r in before]


def test_review_rows_are_highlighted(make_csv, make_template, tmp_path):
    _, out, *_ = _run(make_csv, make_template, tmp_path,
                      review=[False, False, False, False, True])
    ws = openpyxl.load_workbook(out)['RAW']
    assert ws.cell(8, 10).fill.start_color.rgb.endswith('FFFF00')
    assert not ws.cell(3, 10).fill.start_color.rgb.endswith('FFFF00')


def test_rows_already_in_template_are_skipped(make_csv, make_template, tmp_path):
    _, out, txns, _, _ = _run(make_csv, make_template, tmp_path)
    # feed the output back in as the template: nothing new should remain
    again, dropped = prepare_new_transactions(parse_csv(make_csv()), out)
    assert again == [] and dropped == len(txns)


def test_balance_mismatch_is_reported(make_csv, make_template, tmp_path):
    template = make_template()
    txns, _ = prepare_new_transactions(parse_csv(make_csv()), template)
    txns[1]['money_in'] += 10                       # statement no longer adds up
    problems = write_workbook(txns, ['Income'] * len(txns), template, tmp_path / 'o.xlsx')
    assert problems


# ── categorisation that needs no API key ─────────────────────────────────────

def test_client_history_decides_without_the_ai(monkeypatch):
    monkeypatch.delenv('ANTHROPIC_API_KEY', raising=False)
    examples = [{'description': 'ACME OPTICAL SUPPLIES LTD', 'category': 'Purchase',
                 'direction': 'out'}] * 3
    txns = [{'description': 'ACME OPTICAL SUPPLIES LTD 123', 'money_in': 0, 'money_out': 50},
            {'description': 'MYSTERY PAYEE', 'money_in': 0, 'money_out': 9}]
    cats, review = cat.categorise_detailed(txns, client_examples=examples)
    assert cats == ['Purchase', 'Unknown']
    assert review == [False, True]                  # history is trusted, unknown is flagged


def test_category_spellings_follow_the_client():
    canon = cat._make_canonicaliser(['Directors Loan Account'])
    assert canon('DLA') == 'Directors Loan Account'
    assert canon('Bank charge?') == 'Bank charges'
    assert canon('') == 'Unknown'


def test_account_numbers_are_scrubbed_before_sending():
    out = cat._scrub('To A/C 80522920 , 12-34-56 , 12345678 , J SMITH')
    assert '80522920' not in out and '12-34-56' not in out and '12345678' not in out

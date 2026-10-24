"""Regression tests for the findings of the code review (all data invented)."""

from datetime import date, datetime

import openpyxl
import pytest

from core import categorise as cat
from core.columns import resolve_columns
from core.excel_writer import (_capture_pivot_grids, _seed_relocated_pivot_grids,
                               prepare_new_transactions, write_workbook)
from parsers._common import salvage_json_array
from parsers.parse_csv import parse_csv
from parsers.parse_xls import parse_xls


def template(tmp_path, headers, rows, name='t.xlsx', analysis=None):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'RAW'
    ws.append(headers)
    for r in rows:
        ws.append(r)
    if analysis:
        an = wb.create_sheet('Analysis')
        for r in analysis:
            an.append(r)
    p = tmp_path / name
    wb.save(p)
    return p


def raw_rows(n=12, extra=()):
    out = []
    bal = 1000.0
    for i in range(n):
        d = date(2025, 3, 1) + (date(2025, 3, 2) - date(2025, 3, 1)) * i
        bal -= 10
        out.append([i + 1, 'FY24', '24/25', datetime(d.year, d.month, d.day), f'PAYEE {i % 3}',
                    None, 10.0, bal, f'=F{i + 2}-G{i + 2}', 'Purchase', *extra])
    out[-1][7] = 1000.0        # so the invented statement's balances reconcile
    return out


HEADERS = ['No', 'SA', 'ACC', 'Date', 'Description', 'Paid in', 'Paid out',
           'Balance', 'NET', 'UC Category']


def run(make_csv, tmpl, tmp_path):
    txns, _ = prepare_new_transactions(parse_csv(make_csv()), tmpl)
    warnings = []
    out = tmp_path / 'out.xlsx'
    write_workbook(txns, ['Purchase', 'Income', 'Petrol', 'Petrol', 'Parking'], tmpl, out,
                   warnings=warnings)
    return openpyxl.load_workbook(out), txns, warnings


# 1. one stray formula in a history cell must not make the column formula-only
def test_one_old_formula_in_a_category_cell_does_not_hijack_the_column(make_csv, tmp_path):
    rows = raw_rows()
    rows[3][9] = '=J4'                 # "same as above", typed once
    wb, txns, _ = run(make_csv, template(tmp_path, HEADERS, rows), tmp_path)
    ws = wb['RAW']
    new = [ws.cell(r, 10).value for r in range(14, 14 + len(txns))]
    assert new == ['Purchase', 'Income', 'Petrol', 'Petrol', 'Parking']


def test_a_formula_only_on_matching_sign_rows_is_still_replicated(make_csv, tmp_path):
    rows = raw_rows()
    for r in rows:                     # Paid in holds a formula only on money-in rows
        r[5] = None
    rows[1][5], rows[1][6] = '=IF(H3>0,H3,"")', None
    rows[1][7] = rows[0][7] + 20       # fake a credit so the model is consistent
    wb, txns, _ = run(make_csv, template(tmp_path, HEADERS, rows), tmp_path)
    ws = wb['RAW']
    assert ws.cell(15, 6).value == '=IF(H15>0,H15,"")'     # the credit row (2nd new txn)


# 3. the Analysis grid is built on the year column whose labels it uses
def test_analysis_uses_the_year_column_matching_its_headings(make_csv, tmp_path):
    headers = ['No', 'Date', 'Year', 'Personal', 'Description', 'Paid in', 'Paid out',
               'Balance', 'NET', 'UC desc']
    rows = []
    bal = 1000.0
    for i in range(12):
        d = datetime(2025, 3, 1 + i)
        bal -= 10
        rows.append([i + 1, d, 'FY25', '24/25', f'PAYEE {i % 3}', None, 10.0, bal,
                     f'=F{i + 2}-G{i + 2}', 'Purchase'])
    rows[-1][7] = 1000.0
    analysis = [['Sum of Net', 'Column Labels'], ['Row Labels', 'FY25', 'FY26'],
                ['Purchase', -120, None], ['Income', None, None], ['Grand Total', -120, None]]
    tmpl = template(tmp_path, headers, rows, analysis=analysis)
    wb, _, _ = run(make_csv, tmpl, tmp_path)
    f = wb['Analysis']['B3'].value
    assert isinstance(f, str) and f.startswith('=SUMIFS') and "'RAW'!$C:$C" in f, f   # Year, not Personal


def test_analysis_left_alone_when_no_year_column_matches_its_headings(make_csv, tmp_path):
    headers = ['No', 'Date', 'Year', 'Personal', 'Description', 'Paid in', 'Paid out',
               'Balance', 'NET', 'UC desc']
    rows = []
    bal = 1000.0
    for i in range(12):
        bal -= 10
        rows.append([i + 1, datetime(2025, 3, 1 + i), 'FY25', '24/25', f'PAYEE {i % 3}', None,
                     10.0, bal, f'=F{i + 2}-G{i + 2}', 'Purchase'])
    rows[-1][7] = 1000.0
    analysis = [['Sum of Net', 'Column Labels'], ['Row Labels', 'Y/E Mar 25', 'Y/E Mar 26'],
                ['Purchase', -120, None]]
    tmpl = template(tmp_path, headers, rows, analysis=analysis)
    wb, _, _ = run(make_csv, tmpl, tmp_path)
    assert wb['Analysis']['B3'].value == -120           # untouched


# 4 / 7 / 5. header handling
def test_year_columns_found_despite_formatted_empty_rows_below(tmp_path):
    from openpyxl.styles import Font
    tmpl = template(tmp_path, ['No', 'Date', 'Year', 'Personal', 'Description', 'Paid out', 'UC desc'],
                    [[i + 1, datetime(2025, 3, 1 + i), 'FY25', '24/25', 'X', 1.0, 'Y'] for i in range(12)])
    wb = openpyxl.load_workbook(tmpl)
    ws = wb['RAW']
    for r in range(14, 400):
        ws.cell(r, 3).font = Font(bold=True)   # styled but empty: openpyxl counts them in max_row
    wb.save(tmpl)
    rc = resolve_columns(openpyxl.load_workbook(tmpl)['RAW'])
    assert rc['fy'] == 3 and rc['acc'] == 4


def test_transaction_date_header_is_a_date_column(tmp_path):
    ws = openpyxl.Workbook().active
    ws.append(['No', 'Transaction Date', 'Description', 'Paid out', 'UC category'])
    assert resolve_columns(ws)['date'] == 2


def test_punctuation_only_header_is_not_the_number_column(tmp_path):
    ws = openpyxl.Workbook().active
    ws.append(['Date', 'Description', 'Paid out', 'UC category', '?'])
    assert resolve_columns(ws)['no'] is None


# 6. a tick column is never invented on new rows
def test_checked_column_is_not_filled_from_history(make_csv, tmp_path):
    rows = raw_rows(extra=('Y',))
    tmpl = template(tmp_path, HEADERS + ['Checked'], rows)
    wb, txns, _ = run(make_csv, tmpl, tmp_path)
    ws = wb['RAW']
    assert all(ws.cell(r, 11).value is None for r in range(14, 14 + len(txns)))


# 8. no false "left blank" warning for money columns
def test_no_warning_for_paid_in_on_an_expense_only_statement(tmp_path, make_template):
    csv = tmp_path / 's.csv'
    csv.write_text('Date,Description,Paid in,Paid out,Balance\n'
                   + ''.join(f'0{i}/04/2025,SHOP {i},,5.00,{1000 - 5 * i:.2f}\n' for i in range(1, 6)),
                   encoding='utf-8')
    tmpl = make_template()
    txns, _ = prepare_new_transactions(parse_csv(csv), tmpl)
    warnings = []
    write_workbook(txns, ['Lunch'] * len(txns), tmpl, tmp_path / 'o.xlsx', warnings=warnings)
    assert not any('Paid in' in w for w in warnings)


# 10. a short statement crossing a month end with swapped dates
def test_swapped_dates_fixed_on_a_short_month_crossing_statement(tmp_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(['Date', 'Description', 'Amount', 'Balance'])
    for shown, amt, bal in [('25/03/2025', -1, 99), ('28/03/2025', -1, 98), ('31/03/2025', -1, 97),
                            (datetime(2025, 1, 4), -1, 96), (datetime(2025, 5, 4), -1, 95),
                            (datetime(2025, 10, 4), -1, 94)]:
        ws.append([shown, 'X', amt, bal])
    p = tmp_path / 's.xlsx'
    wb.save(p)
    assert [t['date'] for t in parse_xls(p)][3:] == [date(2025, 4, 1), date(2025, 4, 5), date(2025, 4, 10)]


# 11. prose before the array
def test_salvage_ignores_a_bracket_in_prose():
    out = salvage_json_array('Results [1]: [{"category": "Lunch", "sure": true}]')
    assert out == [{'category': 'Lunch', 'sure': True}]


# 12. words that start with a month abbreviation
def test_month_like_words_are_not_stripped_from_payees():
    assert 'JUNCTION' in cat._payee_key('CARD 11 JUNCTION 32 SERVICES')
    assert 'OCTOPUS' in cat._payee_key('DD 12 OCTOPUS 2024')
    assert cat._payee_key('4445 09MAY23 CD , MCDONALDS 1194') == 'MCDONALDS'


# 13. the client's own history beats the HMRC reference rule
def test_client_history_beats_hmrc_reference(monkeypatch):
    monkeypatch.setenv('ANTHROPIC_API_KEY', '')
    history = [{'description': 'HMRC Cumbernauld', 'category': 'HMRC', 'direction': 'out'}] * 4
    t = {'description': 'HMRC Cumbernauld', 'reference': '120PP034375842609', 'money_in': 0, 'money_out': 900}
    cats, review = cat.categorise_detailed([t], client_examples=history)
    assert cats == ['HMRC'] and review == [False]


# short payees
@pytest.mark.parametrize('payee', ['BP', 'EE', 'O2'])
def test_two_character_payees_match_history(monkeypatch, payee):
    monkeypatch.setenv('ANTHROPIC_API_KEY', '')
    history = [{'description': payee, 'category': 'Petrol', 'direction': 'out'}] * 3
    cats, review = cat.categorise_detailed([{'description': payee, 'money_in': 0, 'money_out': 40}],
                                           client_examples=history)
    assert cats == ['Petrol'] and review == [False]


# Howayda: a tax-year column whose new year's first transaction is a few days after 6 April

def _messy_tax_year_pairs(label_style):
    """The client's own rows: a short first group, a second one that runs
    Aug 2023 - 5 Apr 2025 (the accountant's mistake), then the new year
    starting with the first transaction on 9 Apr 2025."""
    from datetime import timedelta
    pairs = []
    d = date(2023, 8, 1)
    while d <= date(2025, 7, 30):
        if d <= date(2023, 8, 8):
            lab = '23/24'
        elif d <= date(2025, 4, 5):
            lab = '24/25'
        elif d >= date(2025, 4, 9):
            lab = '25/26'
        else:
            d += timedelta(days=1)
            continue
        if label_style == 'fy':
            lab = 'FY' + lab[3:]
        pairs.append((d, lab))
        d += timedelta(days=3)
    return pairs


def test_tax_year_column_switches_on_6_april_despite_messy_history():
    from core.year_labels import YearLabeller
    lab = YearLabeller(_messy_tax_year_pairs('acc'))
    assert lab.boundary == (4, 6)
    assert lab.label_for(date(2026, 4, 5)) == '25/26'
    assert lab.label_for(date(2026, 4, 6)) == '26/27'
    assert lab.label_for(date(2026, 4, 8)) == '26/27'


def test_company_year_labels_are_not_snapped_to_6_april():
    from core.year_labels import YearLabeller
    lab = YearLabeller(_messy_tax_year_pairs('fy'))
    assert lab.boundary != (4, 6)             # FY labels keep the learned boundary


def test_personal_column_end_to_end_switches_on_6_april(tmp_path):
    headers = ['No', 'Date', 'Year', 'Personal', 'Description', 'Paid in', 'Paid out',
               'Balance', 'NET', 'UC desc']
    rows, bal = [], 1000.0
    for i, (d, personal) in enumerate(_messy_tax_year_pairs('acc')):
        bal -= 1
        fy = 'FY24' if d < date(2024, 8, 1) else 'FY25'
        rows.append([i + 1, datetime(d.year, d.month, d.day), fy, personal, f'PAYEE {i % 3}', None,
                     1.0, round(bal, 2), f'=F{i + 2}-G{i + 2}', 'Purchase'])
    last_bal = rows[-1][7]
    tmpl = template(tmp_path, headers, rows)
    csv = tmp_path / 's.csv'
    lines = ['Date,Description,Paid in,Paid out,Balance']
    b = last_bal
    for d in (date(2026, 4, 1), date(2026, 4, 3), date(2026, 4, 6), date(2026, 4, 7), date(2026, 4, 8)):
        b = round(b - 2, 2)
        lines.append(f'{d:%d/%m/%Y},SHOP,,2.00,{b:.2f}')
    csv.write_text('\n'.join(lines) + '\n', encoding='utf-8')
    txns, _ = prepare_new_transactions(parse_csv(csv), tmpl)
    out = tmp_path / 'o.xlsx'
    write_workbook(txns, ['Lunch'] * 5, tmpl, out)
    ws = openpyxl.load_workbook(out)['RAW']
    first_new = len(rows) + 2
    assert [ws.cell(first_new + i, 4).value for i in range(5)] == ['25/26', '25/26', '26/27', '26/27', '26/27']
    assert all(ws.cell(first_new + i, 3).value == 'FY26' for i in range(5))   # company year Aug-Jul unchanged


# Howayda: a first-pass column (UC desc) and a reviewer's final column (MUC desc)

def _two_column_template(tmp_path):
    headers = ['No', 'Date', 'Year', 'Description', 'Paid in', 'Paid out', 'Balance',
               'NET', 'MUC desc', 'UC desc']
    rows, bal = [], 1000.0
    for i in range(30):
        bal -= 10
        first = ['DLA', 'Food?', 'Unknown'][i % 3]
        final = ['Dividend', 'Food- client', 'Dividend'][i % 3]    # what the reviewer wrote
        rows.append([i + 1, datetime(2025, 1, 1 + i), 'FY25', f'PAYEE {i % 3}', None, 10.0,
                     round(bal, 2), f'=E{i + 2}-F{i + 2}', final, first])
    rows[-1][6] = 1000.0
    analysis = [['Sum of Net', 'Column Labels'], ['Row Labels', 'FY25', 'FY26'],
                ['Dividend', -200, None], ['Food- client', -100, None], ['Grand Total', -300, None]]
    return template(tmp_path, headers, rows, analysis=analysis)


def test_analysis_is_built_on_the_column_its_labels_match(make_csv, tmp_path):
    wb, _, warnings = run(make_csv, _two_column_template(tmp_path), tmp_path)
    f = wb['Analysis']['B3'].value
    assert isinstance(f, str) and "'RAW'!$I:$I" in f, f          # MUC desc (I), not UC desc (J)
    assert any('reviewer' in w for w in warnings)


def test_client_examples_use_the_reviewed_category(tmp_path):
    from core.excel_writer import extract_client_examples
    ex = extract_client_examples(_two_column_template(tmp_path))
    cats = {e['description']: e['category'] for e in ex}
    assert cats['PAYEE 0'] == 'Dividend' and cats['PAYEE 1'] == 'Food- client'


def test_reviewed_column_stays_empty_on_new_rows(make_csv, tmp_path):
    wb, txns, _ = run(make_csv, _two_column_template(tmp_path), tmp_path)
    ws = wb['RAW']
    for r in range(32, 32 + len(txns)):
        assert ws.cell(r, 9).value is None and ws.cell(r, 10).value


# A second review column with no 'MUC' naming (e.g. 'UC Category' / 'UC Category 1')

def _numbered_two_column_template(tmp_path):
    headers = ['No', 'Date', 'Year', 'Description', 'Paid in', 'Paid out', 'Balance',
               'NET', 'UC Category', 'UC Category 1']
    rows, bal = [], 1000.0
    for i in range(30):
        bal -= 10
        draft = ['DLA', None, None][i % 3]              # mostly blank: only some rows reviewed
        final = ['Dividend', 'Food- client', 'Dividend'][i % 3]   # always filled: the real answer
        rows.append([i + 1, datetime(2025, 1, 1 + i), 'FY25', f'PAYEE {i % 3}', None, 10.0,
                     round(bal, 2), f'=E{i + 2}-F{i + 2}', draft, final])
    rows[-1][6] = 1000.0
    analysis = [['Sum of Net', 'Column Labels'], ['Row Labels', 'FY25', 'FY26'],
                ['Dividend', -200, None], ['Food- client', -100, None], ['Grand Total', -300, None]]
    return template(tmp_path, headers, rows, analysis=analysis)


def test_numbered_second_column_is_recognised_as_reviewed(tmp_path):
    ws = openpyxl.load_workbook(_numbered_two_column_template(tmp_path))['RAW']
    rc = resolve_columns(ws)
    hdr = {c: ws.cell(1, c).value for c in range(1, ws.max_column + 1)}
    assert hdr[rc['uc']] == 'UC Category'
    assert [hdr[c] for c in rc['reviewed']] == ['UC Category 1']


def test_analysis_built_on_the_numbered_reviewed_column(make_csv, tmp_path):
    wb, _, warnings = run(make_csv, _numbered_two_column_template(tmp_path), tmp_path)
    f = wb['Analysis']['B3'].value
    assert isinstance(f, str) and "'RAW'!$J:$J" in f, f       # UC Category 1 (J), not UC Category (I)
    assert any('reviewer' in w for w in warnings)


def test_single_category_column_is_never_mistaken_for_reviewed(make_csv, tmp_path):
    # a normal single-column template must not grow a phantom 'reviewed' column
    ws = openpyxl.load_workbook(template(tmp_path, HEADERS, raw_rows()))['RAW']
    assert resolve_columns(ws)['reviewed'] == []


# A fiscal year whose only recorded start is a few days into a new month
# (one straggler row aside, there is no earlier year to measure a gap from)

def test_single_year_boundary_snaps_to_the_1st_when_start_day_is_early():
    from core.year_labels import YearLabeller
    pairs = [(date(2024, 3, 2) + __import__('datetime').timedelta(days=i), 'FY24') for i in range(0, 360, 3)]
    lab = YearLabeller(pairs)
    assert lab.boundary == (3, 1)


def test_single_year_boundary_not_snapped_when_start_day_is_mid_month():
    from core.year_labels import YearLabeller
    pairs = [(date(2024, 3, 15) + __import__('datetime').timedelta(days=i), 'FY24') for i in range(0, 360, 3)]
    lab = YearLabeller(pairs)
    assert lab.boundary != (3, 1)


# A literal empty-string cell in the template can come back as None after
# openpyxl's save round trip, even on a row write_workbook never touches —
# that must not be reported as "original rows modified".

def test_empty_string_cells_are_not_flagged_as_modified(make_csv, tmp_path):
    rows = raw_rows()
    # a dedicated (if minimal) Analysis sheet, so a small formula-bearing RAW
    # sheet isn't itself mistaken for the Analysis sheet (see _sheet_score)
    tmpl = template(tmp_path, HEADERS, rows, analysis=[['Row Labels', 'FY24']])
    wb = openpyxl.load_workbook(tmpl)
    ws = wb['RAW']
    for r in range(2, 14):
        ws.cell(r, 8).value = ''        # Balance held as a literal '' in the template
    wb.save(tmpl)

    txns, _ = prepare_new_transactions(parse_csv(make_csv()), tmpl)
    problems = write_workbook(txns, ['Purchase', 'Income', 'Petrol', 'Petrol', 'Parking'],
                              tmpl, tmp_path / 'out.xlsx')
    assert problems == []


# A client's real template (a genuine native Excel PivotTable sharing a sheet
# with the Analysis grid) showed the pivot's old spot overwritten by the
# Analysis rebuild and its new, relocated spot sitting completely blank until
# Excel refreshes it on open — which, to a non-technical user, looks exactly
# like "the pivot table isn't working, it's just numbers". The pivot's last
# cached grid must be carried over to its new spot instead of leaving a void.

class _FakeLocation:
    def __init__(self, ref):
        self.ref = ref


class _FakePivot:
    def __init__(self, ref):
        self.location = _FakeLocation(ref)


def test_relocated_pivot_grid_is_seeded_instead_of_left_blank():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Analysis'
    ws['A1'], ws['B1'] = 'Row Labels', 'FY25'
    ws['A2'], ws['B2'] = 'Food', -50.5
    ws['A3'], ws['B3'] = 'Grand Total', -50.5
    ws['B2'].number_format = '#,##0.00'

    piv = _FakePivot('A1:B3')
    ws._pivots = [piv]

    # captured before the Analysis rebuild overwrites the pivot's own cells
    grids = _capture_pivot_grids(ws)
    for r in range(1, 4):
        for c in range(1, 3):
            ws.cell(r, c).value = None
    ws['A1'] = 'Food'   # the SUMIFS grid now lives where the pivot used to be

    # the pivot is relocated clear of the grid ...
    piv.location.ref = 'E1:F3'
    # ... and its last-known grid should reappear at the new spot
    _seed_relocated_pivot_grids(ws, grids)

    assert [ws['E1'].value, ws['F1'].value] == ['Row Labels', 'FY25']
    assert [ws['E2'].value, ws['F2'].value] == ['Food', -50.5]
    assert ws['F2'].number_format == '#,##0.00'
    assert ws['E3'].value == 'Grand Total'
    assert ws['E1'].comment is not None


def test_pivot_already_clear_of_the_grid_is_left_untouched():
    # idempotent: a pivot that was never relocated must not be re-seeded or
    # commented on top of the client's own, already-correct cells
    wb = openpyxl.Workbook()
    ws = wb.active
    ws['A1'] = 'Row Labels'
    piv = _FakePivot('A1:A1')
    ws._pivots = [piv]
    grids = _capture_pivot_grids(ws)
    _seed_relocated_pivot_grids(ws, grids)   # ref unchanged -> no-op
    assert ws['A1'].comment is None
